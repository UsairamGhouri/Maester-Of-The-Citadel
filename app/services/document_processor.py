"""
Document Processing Service
Handles text extraction from various formats: PDF, EPUB, DOCX, PPTX
"""
import os
import logging
from bs4 import BeautifulSoup
import ebooklib
from ebooklib import epub
import docx
from pptx import Presentation
from app.services.pdf_processor import extract_text_by_page, get_pdf_metadata

logger = logging.getLogger('CourseHelper.document_processor')

def get_document_metadata(file_path: str) -> dict:
    """Extract basic metadata from a file based on its extension."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return get_pdf_metadata(file_path)
    
    # Generic metadata for non-PDFs
    return {
        'title': os.path.basename(file_path),
        'author': None,
        'total_pages': 0  # Will be updated during extraction
    }

def extract_text(file_path: str) -> list[dict]:
    """Extract text from a document, returning a list of pages.
    For non-paginated formats (EPUB, DOCX), we create arbitrary "pages".
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return extract_text_by_page(file_path)
    elif ext == '.epub':
        return _extract_epub(file_path)
    elif ext == '.docx':
        return _extract_docx(file_path)
    elif ext == '.pptx':
        return _extract_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

def _extract_epub(file_path: str) -> list[dict]:
    book = epub.read_epub(file_path)
    pages = []
    page_num = 1
    
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            html = item.get_content()
            soup = BeautifulSoup(html, 'html.parser')
            text = soup.get_text(separator='\n', strip=True)
            if text:
                pages.append({'page': page_num, 'text': text})
                page_num += 1
                
    return pages

def _extract_docx(file_path: str) -> list[dict]:
    doc = docx.Document(file_path)
    pages = []
    
    # DOCX doesn't have strict pages, so we'll chunk by every N paragraphs 
    # to simulate pages for our pipeline.
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    chunk_size = 20 # paragraphs per "page"
    for i in range(0, len(paragraphs), chunk_size):
        chunk = "\n".join(paragraphs[i:i+chunk_size])
        pages.append({
            'page': (i // chunk_size) + 1,
            'text': chunk
        })
        
    return pages

def _extract_pptx(file_path: str) -> list[dict]:
    prs = Presentation(file_path)
    pages = []
    
    for i, slide in enumerate(prs.slides):
        text_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    text_parts.append(text)
        
        if text_parts:
            pages.append({
                'page': i + 1,
                'text': "\n".join(text_parts)
            })
            
    return pages
